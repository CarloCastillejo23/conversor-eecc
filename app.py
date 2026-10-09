import streamlit as st
import pypdf
import pandas as pd
import io
import re
import datetime
import traceback
import urllib.request
from PIL import Image as PILImage
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.drawing.image import Image as OpenpyxlImage

# --- VERIFICADOR DE LLAVE MAESTRA ---
try:
    import pikepdf
    HAS_PIKEPDF = True
except ImportError:
    HAS_PIKEPDF = False

st.set_page_config(
    page_title="Convertidor Universal de Estados de Cuenta | Respaldo Tributario",
    page_icon="📈",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# --- ESTILOS VISUALES Y MAQUETACIÓN (CSS PERSONALIZADO) ---
st.markdown("""
<style>
@import url('https://fonts.cdnfonts.com/css/bahnschrift');

html, body, p, h1, h2, h3, h4, h5, h6, div, span, button, input, label, a {
    font-family: 'Bahnschrift', 'Segoe UI', sans-serif;
}
.material-symbols-rounded, .material-icons, [class*="Icon"] {
    font-family: 'Material Symbols Rounded', 'Material Icons', sans-serif !important;
}
#MainMenu, header, footer { visibility: hidden; }
.stApp { background-color: #F8FBFD; }
.block-container { padding-top: 1rem !important; padding-bottom: 0rem !important; max-width: 1140px !important; }

/* NAVBAR */
.nav-container { display: flex; justify-content: space-between; align-items: center; padding: 0.8rem 0rem 1.5rem 0rem; border-bottom: 1px solid #EBF1F6; margin-bottom: 2rem; flex-wrap: wrap; gap: 1rem; }
.nav-left { display: flex; align-items: center; gap: 1.2rem; flex-wrap: wrap; }
.nav-logo { height: 48px; object-fit: contain; }
.nav-tagline { border-left: 1.5px solid #D5DFE7; padding-left: 1.2rem; color: #6C7A89; font-size: 0.95rem; line-height: 1.2; }
.nav-right { display: flex; align-items: center; gap: 0.5rem; color: #2E1E7E; font-weight: 600; font-size: 0.95rem; }

/* SECCIÓN HERO */
.hero-container { display: flex; justify-content: space-between; align-items: center; margin-bottom: 2rem; flex-wrap: wrap; gap: 2rem; }
.hero-badge { color: #66CCA1; font-weight: 700; font-size: 0.82rem; letter-spacing: 0.12em; text-transform: uppercase; margin-bottom: 0.5rem; }
.hero-title { color: #2E1E7E; font-size: 2.7rem; font-weight: 800; line-height: 1.15; margin-bottom: 0.8rem; }
.hero-subtitle { color: #556575; font-size: 1.15rem; max-width: 540px; line-height: 1.45; margin-bottom: 1.5rem; }
.banks-bar { display: flex; align-items: center; gap: 1.8rem; flex-wrap: wrap; }
.bank-logo { height: 24px; object-fit: contain; mix-blend-mode: multiply; }
.doc-icon { width: 48px; height: 48px; object-fit: contain; }

/* TARJETA PRINCIPAL */
.main-card { background: #FFFFFF; border-radius: 16px; padding: 2.2rem 2.8rem; box-shadow: 0 10px 30px rgba(46, 30, 126, 0.05); border: 1px solid #EEF3F8; margin-bottom: 2rem; }

/* PASOS DEL PROCESO */
.steps-container { display: flex; justify-content: space-between; align-items: center; margin-bottom: 2rem; padding: 0 1rem; flex-wrap: wrap; gap: 1rem; }
.step-item { display: flex; align-items: center; gap: 0.75rem; }
.step-num { width: 32px; height: 32px; border-radius: 50%; display: flex; align-items: center; justify-content: center; font-weight: 700; font-size: 0.95rem; }
.step-active { background-color: #2E1E7E; color: #FFFFFF; }
.step-inactive { background-color: #F0F4F8; color: #8C9BAA; }
.step-text { font-size: 0.9rem; font-weight: 600; color: #2E1E7E; line-height: 1.2; max-width: 140px; }
.step-text-muted { color: #8C9BAA; }
.step-line { flex-grow: 1; height: 1px; background-color: #E2EAF1; margin: 0 1rem; min-width: 50px; }

/* DROPZONE STYLING */
[data-testid="stFileUploadDropzone"] { background: #FDFEFF !important; border: 2px dashed #B8D4FC !important; border-radius: 16px !important; padding: 2.5rem 1rem !important; text-align: center !important; transition: all 0.3s ease; }
[data-testid="stFileUploadDropzone"]:hover { border-color: #2E1E7E !important; background: #F8FBFF !important; }

/* BOTONES STREAMLIT */
.stButton>button { background: #2E1E7E !important; color: #FFFFFF !important; font-size: 1.05rem !important; font-weight: 700 !important; padding: 0.75rem 2rem !important; border-radius: 10px !important; border: none !important; box-shadow: 0 4px 14px rgba(46, 30, 126, 0.25) !important; width: 100% !important; transition: all 0.3s ease !important; }
.stButton>button:hover { background: #231666 !important; box-shadow: 0 6px 18px rgba(46, 30, 126, 0.35) !important; transform: translateY(-1px); }
.stDownloadButton>button { background: #66CCA1 !important; color: #2E1E7E !important; font-size: 1.1rem !important; font-weight: 800 !important; border-radius: 10px !important; border: none !important; box-shadow: 0 4px 14px rgba(102, 204, 161, 0.35) !important; width: 100% !important; }
.stDownloadButton>button:hover { background: #55b78f !important; color: #1A104E !important; }

/* ERRORES */
.stAlert { white-space: pre-wrap !important; word-wrap: break-word !important; }

/* TARJETAS INFERIORES DE CARACTERÍSTICAS */
.features-container { display: flex; justify-content: space-between; gap: 1.5rem; background: #F1F8F8; border: 1px solid #E1EEEE; border-radius: 14px; padding: 1.4rem 2.2rem; margin-bottom: 3.5rem; flex-wrap: wrap; }
.feature-item { display: flex; align-items: center; gap: 0.9rem; flex: 1; min-width: 200px; }
.feature-icon-box { width: 44px; height: 44px; border-radius: 10px; display: flex; align-items: center; justify-content: center; background: #FFFFFF; color: #66CCA1; font-size: 1.4rem; box-shadow: 0 2px 6px rgba(0,0,0,0.04); }
.feature-text { font-size: 0.95rem; font-weight: 700; color: #2E1E7E; line-height: 1.25; }

/* FOOTER */
.footer-container { background: #181145; padding: 2.2rem 2rem; border-radius: 18px 18px 0 0; display: flex; justify-content: space-between; align-items: center; flex-wrap: wrap; gap: 1.5rem; color: #A6B4C9; font-size: 0.9rem; margin-top: 2rem; }
.footer-left { display: flex; align-items: center; flex-wrap: wrap; gap: 1.5rem; }
.footer-logo { height: 38px; object-fit: contain; }
.footer-copy { border-left: 1px solid rgba(255,255,255,0.2); padding-left: 1.5rem; color: #CAD5E2; }
.footer-links { display: flex; gap: 1.5rem; flex-wrap: wrap; }
.footer-links a { color: #CAD5E2; text-decoration: none; white-space: nowrap; transition: color 0.2s ease; }
.footer-links a:hover { color: #66CCA1; }
@media (max-width: 768px) { .footer-copy { border-left: none; padding-left: 0; } }
</style>
""", unsafe_allow_html=True)

# --- NAVBAR ---
st.markdown("""
<div class="nav-container">
    <div class="nav-left">
        <img class="nav-logo" src="https://i.postimg.cc/g21gcsnY/3-(2).png" alt="Respaldo Tributario Logo">
        <div class="nav-tagline">Tu aliado en soluciones<br>contables y tributarias</div>
    </div>
    <div class="nav-right">
        <span>🛡️ Seguro y confiable</span>
    </div>
</div>
""", unsafe_allow_html=True)

# --- SECCIÓN HERO ---
st.markdown("""
<div class="hero-container">
    <div>
        <div class="hero-badge">CONVERTIDOR DE ESTADOS DE CUENTA</div>
        <div class="hero-title">Convertidor Universal<br>de Estados de Cuenta</div>
        <div class="hero-subtitle">Convierte tus estados de cuenta en Excel de forma rápida, segura y sin complicaciones.</div>
        <div class="banks-bar">
            <img src="https://i.postimg.cc/rsGdSkRT/bcp-logo.png" class="bank-logo" alt="BCP">
            <img src="https://i.postimg.cc/WzcFF33n/BBVA-2025.png" class="bank-logo" alt="BBVA">
            <img src="https://i.postimg.cc/6qV7dX4t/Interbank-logo-svg.webp" class="bank-logo" alt="Interbank">
            <img src="https://i.postimg.cc/MGyvkwtH/SCOTIABANK.png" class="bank-logo" alt="Scotiabank">
            <img src="https://i.postimg.cc/bvQsXPLd/banbif-logo-png-seeklogo-502109.png" class="bank-logo" alt="BanBif">
        </div>
    </div>
    <div>
        <div style="background: #FFFFFF; padding: 1.8rem 2.4rem; border-radius: 20px; box-shadow: 0 18px 40px rgba(46,30,126,0.08); display: flex; align-items: center; gap: 1.5rem; flex-wrap: wrap; justify-content: center;">
            <div style="background: #FFF1F0; padding: 1.2rem; border-radius: 16px; border: 1px solid #FFCCC7; display: flex; align-items: center; justify-content: center;">
                <img src="https://i.postimg.cc/ZKTn1Xyq/PDF-file-icon-svg.webp" class="doc-icon" alt="PDF">
            </div>
            <div style="color: #2E1E7E; font-size: 2rem; font-weight: bold;">➔</div>
            <div style="background: #E6F7ED; padding: 1.2rem; border-radius: 16px; border: 1px solid #B7EB8F; display: flex; align-items: center; justify-content: center;">
                <img src="https://i.postimg.cc/QCZ99tt0/Microsoft-Office-Excel-(2019-2025)-svg.webp" class="doc-icon" alt="Excel">
            </div>
        </div>
    </div>
</div>
""", unsafe_allow_html=True)

# --- TARJETA PRINCIPAL ---
st.markdown("""
<div class="main-card">
    <div class="steps-container">
        <div class="step-item">
            <div class="step-num step-active">1</div>
            <div class="step-text">Selecciona tu estado de cuenta</div>
        </div>
        <div class="step-line"></div>
        <div class="step-item">
            <div class="step-num step-inactive">2</div>
            <div class="step-text step-text-muted">Ingresa la contraseña<br>(si corresponde)</div>
        </div>
        <div class="step-line"></div>
        <div class="step-item">
            <div class="step-num step-inactive">3</div>
            <div class="step-text step-text-muted">Convierte y descarga<br>tu archivo Excel</div>
        </div>
    </div>
</div>
""", unsafe_allow_html=True)

archivo_pdf = st.file_uploader("Arrastra tu estado de cuenta aquí o haz clic para subir", type=["pdf"], label_visibility="collapsed")
password_pdf = st.text_input("🔒 Contraseña (opcional)", type="password", placeholder="Ingresa la contraseña si tu archivo la requiere")
btn_convertir = st.button("⚙️ Convertir a Excel")

# --- URLs DE LOS LOGOS DE LOS BANCOS ---
logos_bancos_urls = {
    "BCP_CORRIENTE": "https://i.postimg.cc/rsGdSkRT/bcp-logo.png",
    "BCP_AHORROS": "https://i.postimg.cc/rsGdSkRT/bcp-logo.png",
    "BBVA": "https://i.postimg.cc/WzcFF33n/BBVA-2025.png",
    "INTERBANK": "https://i.postimg.cc/6qV7dX4t/Interbank-logo-svg.webp",
    "SCOTIABANK": "https://i.postimg.cc/MGyvkwtH/SCOTIABANK.png",
    "BANBIF": "https://i.postimg.cc/bvQsXPLd/banbif-logo-png-seeklogo-502109.png"
}

# --- MOTOR DE EXTRACCIÓN Y CONVERSIÓN ---
if btn_convertir:
    if archivo_pdf is None:
        st.warning("⚠️ Primero debes seleccionar o arrastrar un archivo PDF.")
    else:
        if not HAS_PIKEPDF:
            st.error("🚨 FALTA INSTALAR LA LLAVE MAESTRA (pikepdf) 🚨\n\nPor favor, ve a tu GitHub, abre el archivo `requirements.txt` y agrega la palabra `pikepdf` en una nueva línea. Luego, reinicia la aplicación haciendo clic en 'Manage App' -> 'Reboot app'.")
            st.stop()
            
        with st.spinner("Desencriptando y aplicando auditoría contable..."):
            try:
                archivo_bytes = archivo_pdf.getvalue()
                
                try:
                    pdf_document = pikepdf.Pdf.open(io.BytesIO(archivo_bytes))
                    pdf_unlocked = io.BytesIO()
                    pdf_document.save(pdf_unlocked)
                    pdf_unlocked.seek(0)
                except pikepdf.PasswordError:
                    if not password_pdf:
                        st.warning("🔒 Este documento tiene contraseña. Escríbela en la casilla de arriba.")
                        st.stop()
                    try:
                        clave_limpia = password_pdf.strip()
                        pdf_document = pikepdf.Pdf.open(io.BytesIO(archivo_bytes), password=clave_limpia)
                        pdf_unlocked = io.BytesIO()
                        pdf_document.save(pdf_unlocked)
                        pdf_unlocked.seek(0)
                    except pikepdf.PasswordError:
                        st.error("❌ La contraseña ingresada es incorrecta.")
                        st.stop()
                
                reader = pypdf.PdfReader(pdf_unlocked)
                all_text_with_coords = []
                texto_completo = ""
                
                for page_num, page in enumerate(reader.pages):
                    try:
                        ext_text = page.extract_text()
                        if ext_text: 
                            texto_completo += ext_text + " "
                            
                        def visitor_extract(text, cm, tm, fontDict, fontSize):
                            try:
                                if tm is not None and isinstance(tm, (list, tuple)) and len(tm) >= 6:
                                    x, y = tm[4], tm[5]
                                    if text and isinstance(text, str) and text.strip():
                                        all_text_with_coords.append((page_num + 1, round(x, 1), round(y, 1), text.strip()))
                            except:
                                pass 
                                
                        page.extract_text(visitor_text=visitor_extract)
                    except Exception as page_err:
                        continue
                    
                lines_by_page_and_y = {}
                for p, x, y, text in all_text_with_coords:
                    key = (p, y)
                    if key not in lines_by_page_and_y:
                        lines_by_page_and_y[key] = []
                    lines_by_page_and_y[key].append((x, text))
                    
                sorted_keys = sorted(lines_by_page_and_y.keys(), key=lambda k: (k[0], -k[1]))
                
                texto_upper = texto_completo.upper()
                banco = "DESCONOCIDO"
                if "ESTADO DE CUENTA CORRIENTE" in texto_upper and "BCP" in texto_upper:
                    banco = "BCP_CORRIENTE"
                elif "ESTADO DE CUENTA DE AHORRO" in texto_upper and "BCP" in texto_upper:
                    banco = "BCP_AHORROS"
                elif "INTERBANK" in texto_upper and "NEGOCIOS" in texto_upper:
                    banco = "INTERBANK"
                elif "SCOTIA" in texto_upper or "20100043140" in texto_upper:
                    banco = "SCOTIABANK"
                elif "BBVA" in texto_upper or "CONTIAHORRO" in texto_upper:
                    banco = "BBVA"
                elif "BANBIF" in texto_upper:
                    banco = "BANBIF"
                    
                st.info(f"🏦 Banco detectado automáticamente: **{banco.replace('_', ' ')}**")
                
                match_year = re.search(r'\b(202\d)\b', texto_completo)
                doc_year = match_year.group(1) if match_year else str(datetime.datetime.now().year)
                
                saldo_inicial_declarado = None
                for key in sorted_keys:
                    items = sorted(lines_by_page_and_y[key], key=lambda i: i[0]) 
                    combined = " ".join([str(i[1]) for i in items])
                    
                    if banco == "SCOTIABANK" and "Saldo Final al" in combined and re.search(r'\d{4}', combined):
                        try: saldo_inicial_declarado = float(combined.split()[-1].replace(',', ''))
                        except: pass
                        break
                    elif banco in ["BBVA", "BCP_AHORROS"] and "SALDO ANTERIOR" in combined:
                        try: saldo_inicial_declarado = float(combined.split()[-1].replace(',', ''))
                        except: pass
                        break
                    elif banco == "INTERBANK":
                        if len(items) > 3 and not re.search(r'[a-zA-Z]', combined) and str(items[0][1]).replace(',', '').replace('.', '').isdigit():
                            try:
                                numeros = re.findall(r'\d{1,3}(?:,\d{3})*\.\d{2}', str(items[0][1]))
                                saldo_inicial_declarado = float(numeros[1].replace(',', '')) if len(numeros) > 1 else float(numeros[0].replace(',', ''))
                            except: pass
                            break
                    elif banco == "BCP_CORRIENTE":
                        if len(items) > 5 and not re.search(r'[a-zA-Z]', combined):
                            try: saldo_inicial_declarado = float(str(items[0][1]).replace(',', ''))
                            except: pass
                            break

                saldo_previo = saldo_inicial_declarado if saldo_inicial_declarado is not None else 0.0
                data_final = []
                
                for key in sorted_keys:
                    items = sorted(lines_by_page_and_y[key], key=lambda i: i[0]) 
                    combined = " ".join([str(i[1]) for i in items])
                    pagina = key[0]
                    
                    fecha, desc, medio, lugar, sucursal, num_op, hora, cargo, abono, itf, saldo = [""]*11
                    es_transaccion = False

                    if banco == "BANBIF":
                        if re.match(r'^\d{2}/\d{2}/\d{2}', combined):
                            for x, t in items:
                                t = str(t)
                                if x < 100: fecha = t
                                elif 150 <= x < 380: desc += t + " "
                                elif 380 <= x < 450: cargo = t.replace(',', '')
                                elif 450 <= x < 510: abono = t.replace(',', '')
                                elif x >= 510: saldo = t.replace(',', '')
                            desc = desc.strip()
                            if "ITF" in desc.upper():
                                itf = cargo if cargo else abono
                                cargo, abono = "", ""
                            numeros = re.findall(r'\b\d{6,10}\b', desc)
                            if numeros: num_op = numeros[-1]
                            es_transaccion = True

                    elif banco == "SCOTIABANK":
                        if re.match(r'^\d{2}/\d{2}\s', combined):
                            tokens = combined.split()
                            fecha = tokens[0] 
                            saldo = tokens[-1]
                            monto_raw = tokens[-2]
                            try:
                                saldo_float = float(saldo.replace(',', ''))
                                monto_float = float(monto_raw.replace(',', ''))
                                if round(saldo_previo - monto_float, 2) == round(saldo_float, 2): cargo = monto_raw
                                elif round(saldo_previo + monto_float, 2) == round(saldo_float, 2): abono = monto_raw
                                saldo_previo = saldo_float
                            except: cargo = monto_raw
                            if len(tokens) >= 3:
                                num_op = tokens[-3] 
                                medio = tokens[2] 
                                desc = " ".join(tokens[3:-3]) 
                            es_transaccion = True

                    elif banco == "INTERBANK":
                        if re.match(r'^\d{2}/\d{2}\s+\d{2}/\d{2}', combined):
                            tokens = combined.split()
                            if len(tokens) >= 3:
                                fecha = tokens[0]
                                saldo = tokens[-1]
                                monto_raw = tokens[-2]
                                cargo = monto_raw.replace('-', '') if '-' in monto_raw else ""
                                abono = monto_raw if '-' not in monto_raw else ""
                                num_op_match = re.search(r'\b\d{7}\b', combined)
                                if num_op_match: num_op = num_op_match.group(0)
                                middle_tokens = tokens[2:-2]
                                if "WEB" in middle_tokens: medio = "WEB"
                                elif "INTERNO" in middle_tokens: medio = "INTERNO"
                                desc_tokens = [t for t in middle_tokens if t != num_op and t != medio]
                                desc = " ".join(desc_tokens)
                                es_transaccion = True

                    elif banco == "BCP_AHORROS":
                        if re.match(r'^\d{2}[A-Z]{3}\s', combined):
                            raw_str = str(items[0][1]) if len(items)==1 else combined
                            tokens = raw_str.split()
                            if len(tokens) >= 2:
                                fecha = tokens[0]
                                monto = tokens[-1]
                                idx = raw_str.rfind(monto)
                                if idx > 60: abono = monto
                                else: cargo = monto
                                inicio_desc = 2 if re.match(r'^\d{2}[A-Z]{3}$', tokens[1]) else 1
                                desc = " ".join(tokens[inicio_desc:-1])
                                es_transaccion = True

                    elif banco == "BBVA":
                        if len(items) > 3 and re.match(r'^\d{2}-\d{2}$', str(items[0][1])):
                            for x, t in items:
                                t = str(t)
                                if x < 60: fecha = t
                                elif 100 < x < 250: desc += t + " "
                                elif 250 < x < 310: sucursal = t 
                                elif 310 < x < 330: medio = t 
                                elif 330 < x < 380: num_op = t 
                                elif 380 < x < 430:
                                    if '-' in t: cargo = t.replace('-', '')
                                    else: abono = t
                                elif 430 < x < 490: itf = t 
                                elif 490 < x < 550: saldo = t
                            desc = desc.strip()
                            es_transaccion = True

                    elif banco == "BCP_CORRIENTE":
                        if re.match(r'^\d{2}-\d{2}\s', combined):
                            tokens = combined.split()
                            if len(tokens) >= 3:
                                fecha = tokens[0]
                                saldo = tokens[-1]
                                monto_raw = tokens[-2]
                                cargo = monto_raw.replace('-', '') if '-' in monto_raw else ""
                                abono = monto_raw if '-' not in monto_raw else ""
                                middle = tokens[1:-2] 
                                medios_conocidos = ["BPI", "POS", "VEN", "INT", "CAJ", "TLC", "BPT"]
                                medio_index = next((i for i, t in enumerate(middle) if t in medios_conocidos), -1)
                                if medio_index != -1:
                                    desc = " ".join(middle[:medio_index])
                                    rest = middle[medio_index+1:]
                                    medio = middle[medio_index]
                                else:
                                    desc = " ".join(middle[:3]) 
                                    rest = middle[3:]
                                for t in rest:
                                    if re.match(r'^\d{3}-\d{3}$', t): lugar = t 
                                    elif re.match(r'^\d{2}:\d{2}$', t): hora = t 
                                    elif re.match(r'^\d{6}$', t): num_op = t 
                                    elif re.match(r'^\d{4}$', t): sucursal = t 
                                    elif len(t) == 6 and t.isalnum(): hora += f" {t}"
                                es_transaccion = True

                    if es_transaccion:
                        data_final.append([pagina, fecha, desc, medio, lugar, sucursal, num_op, hora.strip(), cargo, abono, itf, saldo])
                        
                columnas = ["PAGINA", "FECHA", "DESCRIPCION", "MEDIO", "LUGAR", "SUCURSAL", "NUMERO DE OPERACION", "HORA U ORIGEN", "CARGO", "ABONO", "ITF", "SALDO"]
                
                def formatear_fecha(fecha_str):
                    fecha_str = str(fecha_str).strip().upper()
                    if not fecha_str or re.search(r'\d{4}', fecha_str): return fecha_str
                    match_banbif = re.match(r'^(\d{2})/(\d{2})/(\d{2})$', fecha_str)
                    if match_banbif: return f"{match_banbif.group(1)}/{match_banbif.group(2)}/20{match_banbif.group(3)}"
                    meses = {'ENE':'01', 'FEB':'02', 'MAR':'03', 'ABR':'04', 'MAY':'05', 'JUN':'06', 'JUL':'07', 'AGO':'08', 'SET':'09', 'SEP':'09', 'OCT':'10', 'NOV':'11', 'DIC':'12'}
                    match1 = re.match(r'^(\d{2})[/-](\d{2})$', fecha_str)
                    if match1: return f"{match1.group(1)}/{match1.group(2)}/{doc_year}"
                    match2 = re.match(r'^(\d{2})-?([A-Z]{3})$', fecha_str)
                    if match2: return f"{match2.group(1)}/{meses.get(match2.group(2), '01')}/{doc_year}"
                    return fecha_str

                def limpiar_numero_espacios(val):
                    if not val or str(val).strip() == "": return None
                    try:
                        num = float(str(val).replace(',', '').strip())
                        return num if num != 0 else None
                    except: return None
                
                def parse_float_seguro(val):
                    if not val or str(val).strip() == "": return 0.0
                    try: return float(str(val).replace(',', '').strip())
                    except: return 0.0

                df_final = pd.DataFrame(data_final, columns=columnas)
                
                if len(df_final) > 0:
                    df_final['FECHA'] = df_final['FECHA'].apply(formatear_fecha)
                    
                    if saldo_inicial_declarado is None:
                        s_primero = parse_float_seguro(data_final[0][11])
                        c_primero = parse_float_seguro(data_final[0][8])
                        a_primero = parse_float_seguro(data_final[0][9])
                        i_primero = parse_float_seguro(data_final[0][10])
                        saldo_inicial = s_primero + c_primero + i_primero - a_primero
                    else:
                        saldo_inicial = saldo_inicial_declarado

                    df_final['CARGO'] = df_final['CARGO'].apply(limpiar_numero_espacios)
                    df_final['ABONO'] = df_final['ABONO'].apply(limpiar_numero_espacios)
                    df_final['ITF'] = df_final['ITF'].apply(limpiar_numero_espacios)
                    df_final['SALDO'] = None 
                    
                    fila_saldo_anterior = pd.DataFrame([["", "", "SALDO ANTERIOR", "", "", "", "", "", None, None, None, None]], columns=columnas)
                    df_final = pd.concat([fila_saldo_anterior, df_final], ignore_index=True)

                    st.success(f"✅ ¡Conversión completada! Se procesaron {len(df_final) - 1} transacciones.")
                    st.dataframe(df_final.head(10))
                    
                    # --- CREACIÓN DEL EXCEL Y PEGADO DE LOGO ---
                    buffer = io.BytesIO()
                    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
                        df_final.to_excel(writer, index=False, startrow=8, sheet_name="Estado de Cuenta")
                        worksheet = writer.sheets["Estado de Cuenta"]
                        
                        color_marca = PatternFill(start_color="2E1E7E", end_color="2E1E7E", fill_type="solid")
                        fuente_blanca = Font(name="Bahnschrift", color="FFFFFF", bold=True)
                        fuente_fallback = Font(name="Bahnschrift", color="FFFFFF", bold=True, size=36, italic=True)
                        borde_blanco = Border(
                            left=Side(style='medium', color="FFFFFF"), right=Side(style='medium', color="FFFFFF"),
                            top=Side(style='medium', color="FFFFFF"), bottom=Side(style='medium', color="FFFFFF")
                        )
                        centro = Alignment(horizontal="center", vertical="center")
                        
                        ancho_tabla = len(df_final.columns)
                        for row in range(1, 8):
                            for col in range(1, ancho_tabla + 1):
                                worksheet.cell(row=row, column=col).fill = color_marca
                                
                        # Magia: Descargar e insertar el logo del banco
                        logo_url = logos_bancos_urls.get(banco)
                        if logo_url:
                            try:
                                req = urllib.request.Request(logo_url, headers={'User-Agent': 'Mozilla/5.0'})
                                with urllib.request.urlopen(req) as response:
                                    img_data = io.BytesIO(response.read())
                                
                                # Convertimos a PNG puro en memoria para asegurar que Excel lo lea
                                pil_img = PILImage.open(img_data)
                                png_io = io.BytesIO()
                                pil_img.save(png_io, format="PNG")
                                png_io.seek(0)
                                
                                img_excel = OpenpyxlImage(png_io)
                                # Ajuste de tamaño corporativo para la cabecera
                                target_height = 80
                                aspect_ratio = pil_img.width / pil_img.height
                                img_excel.height = target_height
                                img_excel.width = int(target_height * aspect_ratio)
                                
                                worksheet.add_image(img_excel, 'B2')
                            except Exception as img_err:
                                print(f"Logo no cargó: {img_err}")
                                celda_banco = worksheet.cell(row=4, column=2, value=f"  {banco.replace('_', ' ')}  ")
                                celda_banco.font = fuente_fallback
                        else:
                            celda_banco = worksheet.cell(row=4, column=2, value=f"  {banco.replace('_', ' ')}  ")
                            celda_banco.font = fuente_fallback
                        
                        col_inicio_resumen = max(ancho_tabla - 3, 6)
                        titulos_resumen = ["SALDO DISPONIBLE", "TOTAL CARGOS", "TOTAL ABONOS", "SALDO FINAL"]
                        
                        for i in range(4):
                            c_tit = worksheet.cell(row=3, column=col_inicio_resumen + i, value=titulos_resumen[i])
                            c_tit.font, c_tit.fill, c_tit.border, c_tit.alignment = fuente_blanca, color_marca, borde_blanco, centro
                            
                            c_val = worksheet.cell(row=4, column=col_inicio_resumen + i)
                            c_val.font, c_val.fill, c_val.border, c_val.alignment = fuente_blanca, color_marca, borde_blanco, centro
                            c_val.number_format = '#,##0.00'
                        
                        last_row = 9 + len(df_final)
                        worksheet.cell(row=4, column=9).value = saldo_inicial 
                        worksheet.cell(row=4, column=10).value = f"=SUM(I11:I{last_row}) + SUM(K11:K{last_row})" 
                        worksheet.cell(row=4, column=11).value = f"=SUM(J11:J{last_row})" 
                        worksheet.cell(row=4, column=12).value = "=I4+K4-J4" 
                        
                        for r in range(10, last_row + 1):
                            if r == 10:
                                worksheet.cell(row=r, column=12).value = "=I4" 
                            else:
                                worksheet.cell(row=r, column=12).value = f"=L{r-1}-I{r}+J{r}-K{r}"
                        
                        for r in range(10, last_row + 1):
                            for c in [9, 10, 11, 12]:
                                worksheet.cell(row=r, column=c).number_format = '#,##0.00'
                            
                        for col in worksheet.columns:
                            max_len = max(len(str(cell.value or '')) for cell in col)
                            worksheet.column_dimensions[col[0].column_letter].width = min(max_len + 3, 50)
                    
                    st.download_button(
                        label="📥 Descargar Excel Corporativo (con Logo)",
                        data=buffer.getvalue(),
                        file_name=f"Estado_Cuenta_{banco}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
                else:
                    st.warning("⚠️ No se identificaron transacciones procesables en el archivo.")
            except Exception as e:
                error_details = traceback.format_exc()
                st.error(f"❌ Ocurrió un error general al procesar el archivo:\n\n{error_details}")

# --- SECCIÓN INFERIOR DE CARACTERÍSTICAS ---
st.markdown("""
<div class="features-container">
    <div class="feature-item">
        <div class="feature-icon-box">🛡️</div>
        <div class="feature-text">Proceso seguro<br>y confiable</div>
    </div>
    <div class="feature-item">
        <div class="feature-icon-box">⚡</div>
        <div class="feature-text">Rápido y<br>sin complicaciones</div>
    </div>
    <div class="feature-item">
        <div class="feature-icon-box">📊</div>
        <div class="feature-text">Compatible con<br>formato Excel</div>
    </div>
</div>
""", unsafe_allow_html=True)

# --- FOOTER CORPORATIVO (LOGO FONDO OSCURO) ---
st.markdown(f"""
<div class="footer-container">
    <div class="footer-left">
        <img class="footer-logo" src="https://i.postimg.cc/Xvm8pC8L/RT-COLOR-PRINCIPAL-2-(9)-(1).png" alt="Respaldo Tributario Footer Logo">
        <div class="footer-copy">© {datetime.datetime.now().year} Respaldo Tributario. Todos los derechos reservados.</div>
    </div>
    <div class="footer-links">
        <a href="#">Términos y condiciones</a>
        <a href="#">Política de privacidad</a>
    </div>
</div>
""", unsafe_allow_html=True)
